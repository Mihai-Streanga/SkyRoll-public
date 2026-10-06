"""Checks on the workbook as written to disk, not on what the program believes it wrote.

The showcase's own list, run after every write:

  1. sheets      the three sheets exist, in order, and nothing else
  2. values      no cell holds a formula
  3. copy        Date_initiale is the statement, cell for cell
  4. mirror      every "statement" value in Verificari equals the same cell in Date_initiale
  5. tolerance   a recalculated cell is red exactly when it is outside the element's tolerance
  6. findings    the red cells and the Avertizari_Erori lines are the same set
  7. payout      recalculated payout = net salary - the deductions, re-derived from the written copy
"""
from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook

from .catalog import Catalog
from .comportamente import rounded
from .scrie import FINDING_FILL, SHEETS, recalc_label, statement_label
from .stat import Statement, marca_key, number


def _d(x) -> Decimal:
    return Decimal(str(number(x)))


def validate(path: str, statement: Statement, catalog: Catalog) -> tuple[bool, list[str]]:
    p: list[str] = []
    wb = load_workbook(path, data_only=False)

    if wb.sheetnames != SHEETS:
        return False, [f"sheets: expected {SHEETS}, found {wb.sheetnames}"]

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    p.append(f"values: formula in {ws.title}!{c.coordinate}")
    if p:
        return False, p          # the numbers below would not mean anything

    di = [list(r) for r in wb["Date_initiale"].iter_rows(values_only=True)]
    expected = [statement.headers] + [[r.get(h) for h in statement.headers] for r in statement.rows]
    if di != expected:
        p.append("copy: Date_initiale differs from the statement")
    headers = di[0]
    original = {marca_key(r[headers.index("Marca")]): dict(zip(headers, r)) for r in di[1:]}

    ver = wb["Verificari"]
    vhead = [c.value for c in ver[1]]
    red: set[tuple[str, str]] = set()
    for row in ver.iter_rows(min_row=2):
        marca = marca_key(row[0].value)
        src = original.get(marca)
        if src is None:
            p.append(f"mirror: Verificari has employee {marca!r}, Date_initiale does not")
            continue
        for name, e in catalog.elements.items():
            if statement_label(name) not in vhead:
                continue
            s_cell = row[vhead.index(statement_label(name))]
            r_cell = row[vhead.index(recalc_label(name))]
            if _d(s_cell.value) != _d(src.get(name)):
                p.append(f"mirror: {marca} / {name}: {s_cell.value} vs {src.get(name)}")
            is_red = (r_cell.fill.fgColor.rgb or "").upper() == FINDING_FILL
            outside = abs(_d(s_cell.value) - _d(r_cell.value)) > Decimal(str(e["toleranta"]))
            if is_red != outside:
                p.append(f"tolerance: {marca} / {name}: red={is_red}, outside tolerance={outside}")
            if is_red:
                red.add((marca, name))
            if e["comportament"] == "rest_plata":
                net = e["argumente"]["net"]
                ded = headers[headers.index(net) + 1: headers.index(name)]
                payout = rounded(_d(src.get(net)) - sum((_d(src.get(d)) for d in ded), Decimal(0)),
                                 e["argumente"]["rotunjire"])
                if payout != _d(r_cell.value):
                    p.append(f"payout: {marca}: recalculated {r_cell.value}, written copy gives {payout}")

    listed = {(marca_key(r[0]), r[1]) for r in wb["Avertizari_Erori"].iter_rows(min_row=2, values_only=True)}
    n_lines = wb["Avertizari_Erori"].max_row - 1
    if listed != red or n_lines != len(red):
        p.append(f"findings: {len(red)} red cells, {n_lines} lines in Avertizari_Erori, "
                 f"not the same set ({sorted(red ^ listed)})")
    return not p, p
