"""Reading a payroll statement: one header row, one row per employee, an optional total row."""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl import load_workbook


def marca_key(value) -> str:
    """The employee ID as text: the identity key of a row.

    IDs are text ("A001", "0042"): converting them to numbers loses leading zeros and
    letters, and sorting them numerically reorders them. A spreadsheet may still store
    a numeric ID as a float (19.0); that one is written back as "19".
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def number(value) -> float:
    """A statement cell as a number; an empty cell is 0."""
    if value is None or value == "":
        return 0.0
    return float(value)


@dataclass
class Statement:
    headers: list[str]
    rows: list[dict]                 # one dict per employee, header -> raw cell value

    def column(self, name: str) -> list:
        return [r.get(name) for r in self.rows]


def _is_total(cells: list) -> bool:
    return any(isinstance(c, str) and c.strip().lower().startswith("total") for c in cells)


def read(path: str) -> Statement:
    """Header on row 1, employees from row 2; reading stops at a 'Total...' row."""
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    data = []
    for cells in rows:
        cells = list(cells)
        if all(c is None or c == "" for c in cells):
            continue
        if _is_total(cells):
            break
        data.append(dict(zip(headers, cells)))
    return Statement(headers=headers, rows=data)
