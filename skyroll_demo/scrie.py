"""The output workbook: values only, no formulas, readable in any spreadsheet program.

Date_initiale      the statement, exactly as read
Verificari         per employee and element: the statement's value next to the recalculated one;
                   a recalculated cell outside tolerance is filled red
Avertizari_Erori   one line per finding, with the catalog's message
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .recalcul import Result

SHEETS = ["Date_initiale", "Verificari", "Avertizari_Erori"]
FINDING_FILL = "FFFFC7CE"
FINDINGS_HEADER = ["Marca", "Element", "Statement", "Recalculated", "Difference", "Message"]


def statement_label(element: str) -> str:
    return f"{element} | statement"


def recalc_label(element: str) -> str:
    return f"{element} | recalculated"


def _num(d):
    return int(d) if d == d.to_integral_value() else float(d)


def write(result: Result, path: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    red = PatternFill("solid", start_color=FINDING_FILL, end_color=FINDING_FILL)

    ws = wb.create_sheet(SHEETS[0])
    ws.append(result.statement.headers)
    for row in result.statement.rows:
        ws.append([row.get(h) for h in result.statement.headers])

    ws = wb.create_sheet(SHEETS[1])
    ws.append(["Marca"] + [lbl for e in result.elements for lbl in (statement_label(e), recalc_label(e))])
    by_row: dict[str, dict] = {}
    for c in result.checks:
        by_row.setdefault(c.marca, {})[c.element] = c
    for r, (marca, checks) in enumerate(by_row.items(), start=2):
        ws.cell(r, 1, marca)
        for i, e in enumerate(result.elements):
            c = checks[e]
            ws.cell(r, 2 + 2 * i, _num(c.in_statement))
            cell = ws.cell(r, 3 + 2 * i, _num(c.recalculated))
            if c.is_finding:
                cell.fill = red

    ws = wb.create_sheet(SHEETS[2])
    ws.append(FINDINGS_HEADER)
    for c in result.findings:
        ws.append([c.marca, c.element, _num(c.in_statement), _num(c.recalculated),
                   _num(c.difference), c.message])

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = bold
        ws.freeze_panes = "B2"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
