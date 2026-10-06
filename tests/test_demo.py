from datetime import date

from openpyxl import load_workbook

from skyroll_demo import SIMPLIFIED, recalcul, scrie, sintetic, stat, valideaza
from skyroll_demo.__main__ import main

MONTH = date(2026, 5, 1)


def test_finds_exactly_the_three_planted_mistakes(statement, cat, params):
    found = {(c.marca, c.element) for c in recalcul.run(statement, cat, params, MONTH).findings}
    assert found == {(p.marca, p.element) for p in sintetic.PLANTED}


def test_clean_statement_has_no_findings(clean_statement, cat, params):
    assert recalcul.run(clean_statement, cat, params, MONTH).findings == []


def test_one_mistake_one_finding(clean_statement, cat, params):
    # a wrong tax, carried on by the payroll system into net salary and payout
    row = clean_statement.rows[4]
    row["Impozit"] += 50
    row["Salariu net"] -= 50
    row["Rest plata"] -= 50
    found = recalcul.run(clean_statement, cat, params, MONTH).findings
    assert [(c.marca, c.element) for c in found] == [("A005", "Impozit")]


def test_same_input_same_output(statement, cat, params):
    a = recalcul.run(statement, cat, params, MONTH).checks
    b = recalcul.run(statement, cat, params, MONTH).checks
    assert a == b


def test_read_stops_at_the_total_row(tmp_path, params):
    headers, rows = sintetic.build(params)
    s = stat.read(sintetic.write(str(tmp_path / "s.xlsx"), headers, rows))
    assert len(s.rows) == 12 and s.rows[-1]["Marca"] == "A012"


def _written(tmp_path, statement, cat, params):
    result = recalcul.run(statement, cat, params, MONTH)
    return scrie.write(result, str(tmp_path / "out.xlsx")), result


def test_written_workbook_passes_its_checks(tmp_path, statement, cat, params):
    path, _ = _written(tmp_path, statement, cat, params)
    assert valideaza.validate(path, statement, cat) == (True, [])


def test_written_workbook_has_values_only(tmp_path, statement, cat, params):
    path, _ = _written(tmp_path, statement, cat, params)
    wb = load_workbook(path)
    assert wb.sheetnames == scrie.SHEETS
    assert all(c.data_type != "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def test_validator_catches_a_tampered_recalculation(tmp_path, statement, cat, params):
    path, _ = _written(tmp_path, statement, cat, params)
    wb = load_workbook(path)
    ws = wb["Verificari"]
    col = [c.value for c in ws[1]].index(scrie.recalc_label("Rest plata")) + 1
    ws.cell(2, col).value += 1
    wb.save(path)
    ok, problems = valideaza.validate(path, statement, cat)
    assert not ok
    assert any(p.startswith("payout: A001") for p in problems)


def test_validator_catches_a_finding_missing_from_the_list(tmp_path, statement, cat, params):
    path, _ = _written(tmp_path, statement, cat, params)
    wb = load_workbook(path)
    wb["Avertizari_Erori"].delete_rows(2)
    wb.save(path)
    ok, problems = valideaza.validate(path, statement, cat)
    assert not ok and any(p.startswith("findings:") for p in problems)


def test_validator_catches_a_formula(tmp_path, statement, cat, params):
    path, _ = _written(tmp_path, statement, cat, params)
    wb = load_workbook(path)
    wb["Verificari"]["B2"] = "=1+1"
    wb.save(path)
    ok, problems = valideaza.validate(path, statement, cat)
    assert not ok and any(p.startswith("values: formula in Verificari!B2") for p in problems)


def test_cli_demo(tmp_path, capsys):
    assert main(["--out", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "Planted: 3 · found: 3 · missed: 0 · unexpected: 0" in out
    assert "Written-file checks: all 7 passed" in out
    assert SIMPLIFIED in out


def test_cli_refuses_and_calculates_nothing(tmp_path, capsys):
    assert main(["--out", str(tmp_path), "--drop-column", "Ore Lucrate"]) == 2
    assert "Guard: REFUSED" in capsys.readouterr().out
    assert not (tmp_path / "SkyRoll_demo.xlsx").exists()
